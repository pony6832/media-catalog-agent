from __future__ import annotations

import os
from pathlib import Path
from typing import Iterable
import uuid
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import MediaRecord, Status


CATALOG_HEADERS = (
    "狀態",
    "檔名",
    "完整路徑",
    "媒體類型",
    "內容描述",
    "重點",
    "關鍵字",
    "拍攝時間",
    "處理時間",
    "Markdown 路徑",
    "備份路徑",
    "錯誤原因",
)

_STATUS_LABELS = {
    Status.PENDING: "待處理",
    Status.PROCESSING: "處理中",
    Status.ANALYZED: "待確認",
    Status.COMPLETED: "完成",
    Status.SKIPPED: "略過",
    Status.FAILED: "失敗",
}

_MEDIA_LABELS = {
    "image/jpeg": "照片 (JPEG)",
    "image/png": "照片 (PNG)",
    "image/webp": "照片 (WebP)",
    "image/heic": "照片 (HEIC)",
    "video/mp4": "影片 (MP4)",
    "video/quicktime": "影片 (MOV)",
    "video/x-matroska": "影片 (MKV)",
    "video/webm": "影片 (WebM)",
}


class ReviewedPathsError(RuntimeError):
    pass


def read_reviewed_paths(excel_path: Path) -> set[str]:
    source = Path(excel_path)
    if not source.is_file():
        return set()
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException):
        return set()
    try:
        sheet = workbook["媒體清冊"]
        headers = {
            cell.value: index
            for index, cell in enumerate(sheet[1], start=1)
            if isinstance(cell.value, str)
        }
        status_column = headers.get("狀態")
        path_column = headers.get("完整路徑")
        if status_column is None or path_column is None:
            return set()
        return {
            str(path_value)
            for status_value, path_value in (
                (
                    sheet.cell(row, status_column).value,
                    sheet.cell(row, path_column).value,
                )
                for row in range(2, sheet.max_row + 1)
            )
            if status_value == "已審核" and path_value
        }
    finally:
        workbook.close()


def read_reviewed_paths_strict(excel_path: Path) -> set[str]:
    source = Path(excel_path)
    if not source.is_file():
        raise ReviewedPathsError("找不到 Excel 媒體清冊")
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException) as error:
        raise ReviewedPathsError("Excel 媒體清冊無法讀取") from error
    try:
        if "媒體清冊" not in workbook.sheetnames:
            raise ReviewedPathsError("Excel 缺少媒體清冊工作表")
        sheet = workbook["媒體清冊"]
        headers = {
            cell.value: index
            for index, cell in enumerate(sheet[1], start=1)
            if isinstance(cell.value, str)
        }
        status_column = headers.get("狀態")
        path_column = headers.get("完整路徑")
        if status_column is None or path_column is None:
            raise ReviewedPathsError("Excel 缺少狀態或完整路徑欄位")
        return {
            str(sheet.cell(row, path_column).value)
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, status_column).value == "已審核"
            and sheet.cell(row, path_column).value
        }
    finally:
        workbook.close()


def write_excel(records: Iterable[MediaRecord], output_path: Path) -> Path:
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    catalog_records = list(records)
    reviewed_paths = read_reviewed_paths(destination)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "媒體清冊"
    sheet.append(CATALOG_HEADERS)

    for record in catalog_records:
        source_path = record.path.resolve()
        path_text = str(source_path)
        sheet.append(
            (
                (
                    "已審核"
                    if path_text in reviewed_paths
                    else _STATUS_LABELS[record.status]
                ),
                record.path.name,
                path_text,
                _MEDIA_LABELS.get(record.media_type, record.media_type),
                record.description,
                "；".join(record.highlights),
                "、".join(record.keywords),
                None,
                (
                    record.updated_at
                    if record.status in {Status.ANALYZED, Status.COMPLETED}
                    else None
                ),
                str(record.markdown_path) if record.markdown_path else None,
                str(record.backup_path) if record.backup_path else None,
                record.error,
            )
        )
        path_cell = sheet.cell(sheet.max_row, 3)
        if source_path.is_file():
            path_cell.hyperlink = source_path.as_uri()
            path_cell.style = "Hyperlink"

    if catalog_records:
        review_validation = DataValidation(
            type="list",
            formula1='"待確認,已審核"',
            allow_blank=False,
            showDropDown=False,
        )
        review_validation.errorTitle = "無效狀態"
        review_validation.error = "請從下拉選單選擇待確認或已審核。"
        review_validation.showErrorMessage = True
        sheet.add_data_validation(review_validation)
        review_validation.add(f"A2:A{len(catalog_records) + 1}")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    widths = (12, 28, 64, 18, 44, 36, 30, 20, 24, 54, 54, 40)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    temporary = destination.with_name(
        f".{destination.stem}.{uuid.uuid4().hex}.tmp{destination.suffix}"
    )
    try:
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, destination)
    finally:
        workbook.close()
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass
    return destination
