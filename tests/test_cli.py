import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

import media_catalog.cli as cli_module
from media_catalog.analysis_mode import AnalysisMode
from media_catalog.batch_analysis import BatchAnalysisResult
from media_catalog.bootstrap import bootstrap_workspace
from media_catalog.cli import main
from media_catalog.database import CatalogDatabase
from media_catalog.inference import Analysis, AnalysisError
from media_catalog.models import Status


class SuccessfulAnalyzer:
    def analyze(self, _: Path) -> Analysis:
        return Analysis("賀卡預覽", ("紅色",), ("賀卡",))


class FailingAnalyzer:
    def analyze(self, _: Path) -> Analysis:
        raise AnalysisError("model stopped")


def _catalog_root_with_one_pending_photo(tmp_path: Path) -> Path:
    root = tmp_path / "media"
    root.mkdir()
    (root / "photo.jpg").write_bytes(b"photo")
    bootstrap_workspace(root)
    return root


def test_cli_start_prints_machine_readable_ready_marker(
    tmp_path: Path, capsys
) -> None:
    root = tmp_path / "media"
    root.mkdir()
    (root / "clip.mp4").write_bytes(b"video")

    exit_code = main(["start", str(root)])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "MEDIA_CATALOG_READY" in output
    assert "added=1" in output
    assert str(root / "媒體整理成果" / "媒體清冊.xlsx") in output


def test_cli_reports_invalid_root_without_creating_fallback(
    tmp_path: Path, capsys
) -> None:
    missing = tmp_path / "missing"

    exit_code = main(["start", str(missing)])

    captured = capsys.readouterr()
    assert exit_code == 2
    assert captured.out == ""
    assert "MEDIA_CATALOG_ERROR" in captured.err
    assert not missing.exists()


def test_cli_analyze_all_prints_fixed_summary(tmp_path: Path, capsys) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)

    exit_code = main(
        [
            "analyze-all",
            str(root),
            "--skill-root",
            str(tmp_path / "media-inventory"),
        ],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert captured.err == ""
    assert captured.out.splitlines()[-1].startswith(
        "MEDIA_ANALYSIS_READY mode=auto analyzed=1 failed=0 skipped=0 remaining=0"
    )
    assert f"catalog={root / '媒體整理成果' / '媒體清冊.xlsx'}" in captured.out


def test_force_cli_passes_mode_and_run_id_to_batch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)
    captured: dict[str, object] = {}

    def fake_batch(*_args, **kwargs):
        captured.update(kwargs)
        return BatchAnalysisResult(1, 0, 0, 0)

    monkeypatch.setattr(cli_module, "analyze_pending", fake_batch)
    monkeypatch.setattr(
        cli_module,
        "read_reviewed_paths",
        lambda _path: {"C:\\reviewed.jpg"},
    )

    exit_code = main(
        [
            "analyze-all",
            str(root),
            "--skill-root",
            str(tmp_path),
            "--mode",
            "force-gemini",
            "--run-id",
            "force-root-1",
        ],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    assert exit_code == 0
    assert captured["mode"] is AnalysisMode.FORCE_GEMINI
    assert captured["run_id"] == "force-root-1"
    assert captured["reviewed_paths"] == {"C:\\reviewed.jpg"}


def test_cli_analyze_all_recovers_interrupted_and_failed_rows(
    tmp_path: Path, capsys
) -> None:
    root = tmp_path / "media"
    root.mkdir()
    (root / "interrupted.jpg").write_bytes(b"one")
    (root / "failed.jpg").write_bytes(b"two")
    workspace = bootstrap_workspace(root).workspace
    database = CatalogDatabase(workspace.database_path)
    records = database.list_records()
    database.set_status(records[0].id, Status.PROCESSING)
    database.set_status(records[1].id, Status.FAILED, error="old failure")

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    captured = capsys.readouterr()
    assert exit_code == 0
    assert "recovered_processing=1" in captured.out
    assert "retried_failed=1" in captured.out
    assert all(
        record.status is Status.ANALYZED
        for record in database.list_records()
    )


def test_cli_analyze_all_does_not_report_ready_with_blank_failed_rows(
    tmp_path: Path, capsys
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=lambda **_kwargs: FailingAnalyzer(),
    )

    captured = capsys.readouterr()
    assert exit_code == 3
    assert "MEDIA_ANALYSIS_INCOMPLETE" in captured.out
    assert "failed=1" in captured.out
    assert "remaining=1" in captured.out


def test_cli_analyze_all_streams_progress_for_slow_computers(
    tmp_path: Path, capsys
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    assert exit_code == 0
    lines = capsys.readouterr().out.splitlines()
    assert lines[0].startswith("MEDIA_ANALYSIS_PROGRESS completed=1 total=1")
    assert "status=analyzed" in lines[0]
    assert lines[-1].startswith("MEDIA_ANALYSIS_READY")


def test_cli_escapes_filename_unsupported_by_console_encoding(
    tmp_path: Path, monkeypatch
) -> None:
    root = tmp_path / "media"
    root.mkdir()
    (root / "视频.jpg").write_bytes(b"photo")
    bootstrap_workspace(root)
    raw = io.BytesIO()
    cp950 = io.TextIOWrapper(raw, encoding="cp950", errors="strict")
    monkeypatch.setattr(sys, "stdout", cp950)

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    cp950.flush()
    output = raw.getvalue().decode("cp950")
    cp950.detach()
    assert exit_code == 0
    assert r"item=\u89c6\u9891.jpg" in output


def test_cli_analyze_all_rebuilds_stale_excel_from_completed_database(
    tmp_path: Path, capsys
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)
    workspace = bootstrap_workspace(root).workspace
    database = CatalogDatabase(workspace.database_path)
    record = database.list_records()[0]
    database.save_analysis(
        record.id,
        description="資料庫完整描述",
        highlights=("完整重點",),
        keywords=("完整關鍵字",),
    )
    workbook = load_workbook(workspace.excel_path)
    sheet = workbook.active
    sheet.cell(2, 1).value = "處理中"
    for column in (5, 6, 7):
        sheet.cell(2, column).value = None
    workbook.save(workspace.excel_path)
    workbook.close()

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    assert exit_code == 0
    assert "MEDIA_ANALYSIS_READY" in capsys.readouterr().out
    workbook = load_workbook(workspace.excel_path, read_only=True)
    row = [workbook.active.cell(2, column).value for column in range(1, 8)]
    workbook.close()
    assert row[0] == "待確認"
    assert row[4:7] == ["資料庫完整描述", "完整重點", "完整關鍵字"]


def test_cli_does_not_repeat_excel_write_after_batch_sync(
    tmp_path: Path, monkeypatch
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)

    def duplicate_write(*_args, **_kwargs):
        raise PermissionError("duplicate final Excel write")

    monkeypatch.setattr(cli_module, "write_excel", duplicate_write)

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=lambda **_kwargs: SuccessfulAnalyzer(),
    )

    assert exit_code == 0


@pytest.mark.parametrize("command", ["start", "resume-processing", "retry-failed"])
def test_cli_mutating_commands_respect_analysis_lock(
    tmp_path: Path, capsys, command: str
) -> None:
    from media_catalog.run_lock import analysis_run_lock

    root = _catalog_root_with_one_pending_photo(tmp_path)
    workspace = bootstrap_workspace(root).workspace

    with analysis_run_lock(workspace.result_root / ".analysis.lock"):
        exit_code = main([command, str(root)])

    assert exit_code == 2
    assert "已有分析程序" in capsys.readouterr().err


def test_cli_reports_runtime_preflight_failure(tmp_path: Path, capsys) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)

    def unavailable(**_kwargs):
        from media_catalog.analysis_runtime import RuntimePreflightError

        raise RuntimePreflightError("找不到模型")

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=unavailable,
    )

    captured = capsys.readouterr()
    assert exit_code == 2
    assert captured.out == ""
    assert "MEDIA_ANALYSIS_ERROR 找不到模型" in captured.err


def test_cli_preflight_failure_preserves_recoverable_rows(
    tmp_path: Path, capsys
) -> None:
    root = tmp_path / "media"
    root.mkdir()
    (root / "interrupted.jpg").write_bytes(b"one")
    (root / "failed.jpg").write_bytes(b"two")
    workspace = bootstrap_workspace(root).workspace
    database = CatalogDatabase(workspace.database_path)
    records = database.list_records()
    database.set_status(records[0].id, Status.PROCESSING)
    database.set_status(records[1].id, Status.FAILED, error="old failure")

    def unavailable(**_kwargs):
        from media_catalog.analysis_runtime import RuntimePreflightError

        raise RuntimePreflightError("FFmpeg 無法使用")

    exit_code = main(
        ["analyze-all", str(root), "--skill-root", str(tmp_path)],
        runtime_builder=unavailable,
    )

    assert exit_code == 2
    assert [record.status for record in database.list_records()] == [
        Status.PROCESSING,
        Status.FAILED,
    ]
    assert "MEDIA_ANALYSIS_ERROR FFmpeg 無法使用" in capsys.readouterr().err


def test_cli_resume_processing_requeues_only_interrupted_items(
    tmp_path: Path, capsys
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)
    workspace = bootstrap_workspace(root).workspace
    database = CatalogDatabase(workspace.database_path)
    record = database.list_records()[0]
    database.set_status(record.id, Status.PROCESSING)

    exit_code = main(["resume-processing", str(root)])

    assert exit_code == 0
    assert "MEDIA_ANALYSIS_RESUMED count=1" in capsys.readouterr().out
    assert database.get_record(record.id).status is Status.PENDING
    workbook = load_workbook(workspace.excel_path, read_only=True)
    assert workbook.active.cell(2, 1).value == "待處理"
    workbook.close()


def test_cli_retry_failed_requeues_failed_items(tmp_path: Path, capsys) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)
    workspace = bootstrap_workspace(root).workspace
    database = CatalogDatabase(workspace.database_path)
    record = database.list_records()[0]
    database.set_status(record.id, Status.FAILED, error="invalid JSON")

    exit_code = main(["retry-failed", str(root)])

    assert exit_code == 0
    assert "MEDIA_ANALYSIS_RETRY_QUEUED count=1" in capsys.readouterr().out
    assert database.get_record(record.id).status is Status.PENDING


def test_cli_verify_sources_reports_success_and_mismatch(
    tmp_path: Path, capsys
) -> None:
    root = _catalog_root_with_one_pending_photo(tmp_path)

    assert main(["verify-sources", str(root)]) == 0
    first = capsys.readouterr()
    assert first.out == "MEDIA_SOURCES_VERIFIED total=1\n"
    (root / "photo.jpg").write_bytes(b"changed")

    assert main(["verify-sources", str(root)]) == 2
    second = capsys.readouterr()
    assert second.out == ""
    assert "MEDIA_ANALYSIS_ERROR" in second.err
