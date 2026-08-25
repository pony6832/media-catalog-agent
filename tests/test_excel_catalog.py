import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from media_catalog.database import CatalogDatabase
from media_catalog.excel_catalog import (
    CATALOG_HEADERS,
    read_reviewed_paths,
    write_excel,
)


def test_write_excel_uses_spec_columns_and_persisted_values(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    photo_path = tmp_path / "旅行照片.jpg"
    photo_path.write_bytes(b"photo")
    record = database.upsert_discovered(photo_path, "abc", "image/jpeg")
    output = tmp_path / "媒體清冊.xlsx"

    saved_path = write_excel([record], output)

    workbook = load_workbook(saved_path, read_only=True)
    sheet = workbook["媒體清冊"]
    assert [cell.value for cell in sheet[1]] == list(CATALOG_HEADERS)
    assert sheet.cell(2, 1).value == "待處理"
    assert sheet.cell(2, 2).value == "旅行照片.jpg"
    assert sheet.cell(2, 3).value == str(photo_path.resolve())
    assert sheet.cell(2, 4).value == "照片 (JPEG)"
    workbook.close()


def test_write_excel_labels_analyzed_records_for_review(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    photo_path = tmp_path / "賀卡.png"
    photo_path.write_bytes(b"photo")
    record = database.upsert_discovered(photo_path, "abc", "image/png")
    analyzed = database.save_analysis(
        record.id,
        description="紅色馬年賀卡。",
        highlights=("金色馬",),
        keywords=("賀卡", "馬年"),
    )

    saved_path = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(saved_path, read_only=True)
    sheet = workbook["媒體清冊"]
    assert sheet.cell(2, 1).value == "待確認"
    assert sheet.cell(2, 5).value == "紅色馬年賀卡。"
    assert sheet.cell(2, 9).value is not None
    workbook.close()


def test_write_excel_adds_review_dropdown_to_status_rows(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    first_path = tmp_path / "first.png"
    second_path = tmp_path / "second.png"
    first_path.write_bytes(b"first")
    second_path.write_bytes(b"second")
    first = database.upsert_discovered(first_path, "first", "image/png")
    second = database.upsert_discovered(second_path, "second", "image/png")

    saved_path = write_excel([first, second], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(saved_path)
    sheet = workbook["媒體清冊"]
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    validation = validations[0]
    assert validation.type == "list"
    assert validation.formula1 == '"待確認,已審核"'
    assert str(validation.sqref) == "A2:A3"
    assert validation.showDropDown is False
    workbook.close()


def test_write_excel_omits_review_dropdown_for_empty_catalog(
    tmp_path: Path,
) -> None:
    saved_path = write_excel([], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(saved_path)
    sheet = workbook["媒體清冊"]
    assert list(sheet.data_validations.dataValidation) == []
    assert sheet.max_row == 1
    workbook.close()


def test_write_excel_preserves_existing_workbook_when_atomic_replace_fails(
    tmp_path: Path, monkeypatch
) -> None:
    output = tmp_path / "媒體清冊.xlsx"
    output.write_bytes(b"existing-workbook")

    def locked_replace(*_args, **_kwargs):
        raise PermissionError("workbook is open")

    monkeypatch.setattr(os, "replace", locked_replace)

    with pytest.raises(PermissionError, match="open"):
        write_excel([], output)

    assert output.read_bytes() == b"existing-workbook"
    assert list(tmp_path.glob(".媒體清冊.*.tmp.xlsx")) == []


def test_path_cell_links_to_existing_source(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "片段 01.mp4"
    source.write_bytes(b"video")
    record = database.upsert_discovered(source, "video-1", "video/mp4")

    output = write_excel([record], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(output)
    path_cell = workbook["媒體清冊"].cell(2, 3)
    assert path_cell.value == str(source.resolve())
    assert path_cell.hyperlink is not None
    assert path_cell.hyperlink.target == source.resolve().as_uri()
    assert path_cell.style == "Hyperlink"
    workbook.close()


def test_path_cell_stays_plain_when_source_is_missing(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "missing.mp4"
    record = database.upsert_discovered(source, "missing", "video/mp4")

    output = write_excel([record], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(output)
    path_cell = workbook["媒體清冊"].cell(2, 3)
    assert path_cell.value == str(source.resolve())
    assert path_cell.hyperlink is None
    assert path_cell.style != "Hyperlink"
    workbook.close()


def test_atomic_rebuild_preserves_reviewed_status_by_path(
    tmp_path: Path,
) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "reviewed.png"
    source.write_bytes(b"photo")
    record = database.upsert_discovered(source, "reviewed", "image/png")
    analyzed = database.save_analysis(
        record.id,
        description="已完成描述",
        highlights=("清晰",),
        keywords=("照片",),
    )
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    workbook = load_workbook(output)
    workbook["媒體清冊"].cell(2, 1).value = "已審核"
    workbook.save(output)
    workbook.close()

    write_excel([analyzed], output)

    workbook = load_workbook(output)
    assert workbook["媒體清冊"].cell(2, 1).value == "已審核"
    workbook.close()
    assert read_reviewed_paths(output) == {str(source.resolve())}


def test_gemini_warning_keeps_analysis_visible_for_manual_review(
    tmp_path: Path,
) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "photo.jpg"
    source.write_bytes(b"photo")
    record = database.upsert_discovered(source, "photo", "image/jpeg")
    analyzed = database.save_analysis(
        record.id,
        description="本地分析仍可使用。",
        highlights=("保留重點",),
        keywords=("本地備援",),
        warning="Gemini 強化失敗:GeminiError",
    )

    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(output, read_only=True)
    row = tuple(
        cell.value for cell in next(workbook["媒體清冊"].iter_rows(min_row=2))
    )
    workbook.close()
    assert row[0] == "待確認"
    assert row[4:7] == (
        "本地分析仍可使用。",
        "保留重點",
        "本地備援",
    )
    assert row[11] == "Gemini 強化失敗:GeminiError"
