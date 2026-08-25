from collections.abc import Sequence
from pathlib import Path

import pytest
from openpyxl import load_workbook

from media_catalog.batch_analysis import BatchAnalysisResult, analyze_pending
from media_catalog.analysis_mode import AnalysisMode
from media_catalog.database import CatalogDatabase
from media_catalog.force_gemini import ForceImageResult
from media_catalog.inference import Analysis, AnalysisError
from media_catalog.models import Status
from media_catalog.run_state import RunStateStore
from media_catalog.scanner import scan
from media_catalog.segment_pipeline import SafeStopRequested, VideoAnalysisResult
from media_catalog.source_guard import capture_source
from media_catalog.workspace import MediaWorkspace


class PathAwareAnalyzer:
    def __init__(self, fail_name: str | None = None) -> None:
        self.fail_name = fail_name
        self.sources: list[Path] = []

    def analyze(self, source: Path) -> Analysis:
        self.sources.append(source)
        if source.name == self.fail_name:
            raise AnalysisError("invalid image")
        return Analysis("賀卡預覽", ("紅色",), ("賀卡",))


class RecordingForceImages:
    def __init__(self, result: ForceImageResult) -> None:
        self.result = result
        self.sources: list[Path] = []

    def analyze(self, source: Path) -> ForceImageResult:
        self.sources.append(source.resolve())
        return self.result


class ForceRuntime:
    def __init__(
        self, store: RunStateStore, force_images: RecordingForceImages
    ) -> None:
        self.run_state = store
        self.segment_pipeline = object()
        self.force_image_analyzer = force_images

    def analyze(self, _source: Path) -> Analysis:
        raise AssertionError("auto photo analysis must not run in force mode")


LOCAL = Analysis("本地描述完整。", ("本地重點",), ("本地",))


def _workspace_with_media(
    tmp_path: Path, names: Sequence[str]
) -> MediaWorkspace:
    root = tmp_path / "media"
    root.mkdir()
    for name in names:
        (root / name).write_bytes(name.encode("utf-8"))
    workspace = MediaWorkspace.from_root(root)
    workspace.ensure_directories()
    database = CatalogDatabase(workspace.database_path)
    scan(workspace.root, database, excluded_roots=(workspace.result_root,))
    return workspace


def test_analyze_pending_continues_after_one_item_fails_and_updates_excel(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("bad.jpg", "good.jpg"))

    result = analyze_pending(
        workspace, PathAwareAnalyzer(fail_name="bad.jpg")
    )

    database = CatalogDatabase(workspace.database_path)
    records = {record.path.name: record for record in database.list_records()}
    assert result == BatchAnalysisResult(
        analyzed=1, failed=1, skipped=0, remaining=1
    )
    assert records["bad.jpg"].status is Status.FAILED
    assert records["good.jpg"].status is Status.ANALYZED
    workbook = load_workbook(workspace.excel_path, read_only=True)
    statuses = [
        row[0]
        for row in workbook.active.iter_rows(min_row=2, values_only=True)
    ]
    workbook.close()
    assert statuses == ["失敗", "待確認"]


def test_analyze_pending_skips_analyzed_and_failed_records_on_rerun(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("bad.jpg", "good.jpg"))
    first_analyzer = PathAwareAnalyzer(fail_name="bad.jpg")
    first = analyze_pending(workspace, first_analyzer)
    second_analyzer = PathAwareAnalyzer()

    second = analyze_pending(workspace, second_analyzer)

    assert first.analyzed == 1
    assert second == BatchAnalysisResult(
        analyzed=0, failed=1, skipped=2, remaining=1
    )
    assert second_analyzer.sources == []


def test_analyze_pending_reports_progress_after_each_durable_checkpoint(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("first.jpg", "second.jpg"))
    events: list[tuple[int, int, str, Status]] = []

    analyze_pending(
        workspace,
        PathAwareAnalyzer(),
        progress=lambda completed, total, record: events.append(
            (completed, total, record.path.name, record.status)
        ),
    )

    assert events == [
        (1, 2, "first.jpg", Status.ANALYZED),
        (2, 2, "second.jpg", Status.ANALYZED),
    ]


def test_analyze_pending_preserves_source_snapshot(tmp_path: Path) -> None:
    workspace = _workspace_with_media(tmp_path, ("good.jpg",))
    source = workspace.root / "good.jpg"
    before = capture_source(source)

    analyze_pending(workspace, PathAwareAnalyzer())

    assert capture_source(source) == before


def test_excel_failure_stops_before_analyzing_the_next_item(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("first.jpg", "second.jpg"))
    analyzer = PathAwareAnalyzer()

    def locked_excel(*_args, **_kwargs):
        raise PermissionError("workbook is locked")

    with pytest.raises(PermissionError, match="locked"):
        analyze_pending(workspace, analyzer, excel_writer=locked_excel)

    statuses = [
        record.status
        for record in CatalogDatabase(workspace.database_path).list_records()
    ]
    assert statuses == [Status.PROCESSING, Status.PENDING]
    assert analyzer.sources == []


def test_a_plus_runtime_continues_sqlite_analysis_when_excel_is_locked(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("clip.mp4", "photo.jpg"))
    store = RunStateStore(
        workspace.database_path, excel_path=workspace.excel_path
    )

    class FakeVideoPipeline:
        def analyze_video(self, record, run_id, *, mode=AnalysisMode.AUTO):
            return VideoAnalysisResult(
                description="影片片段分析完成。",
                highlights=("片段",),
                keywords=("影片",),
                gemini_segments=0,
                needs_review_segments=0,
                failed_segments=0,
            )

    class APlusRuntime(PathAwareAnalyzer):
        run_state = store
        segment_pipeline = FakeVideoPipeline()

    def locked_excel(*_args, **_kwargs):
        raise PermissionError("workbook is locked")

    result = analyze_pending(
        workspace, APlusRuntime(), excel_writer=locked_excel
    )

    records = CatalogDatabase(workspace.database_path).list_records()
    assert all(record.status is Status.ANALYZED for record in records)
    assert result.analyzed == 2
    assert result.remaining == 0
    assert result.excel_sync_pending is True
    run = store.get_run(store.run_id_for_root(workspace.root))
    assert run is not None and run.excel_sync_pending is True
    assert run.status == "incomplete"


def test_safe_stop_returns_current_video_to_pending_without_failure(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("clip.mp4",))
    store = RunStateStore(
        workspace.database_path, excel_path=workspace.excel_path
    )

    class StoppingPipeline:
        def analyze_video(self, _record, run_id, *, mode=AnalysisMode.AUTO):
            store.request_stop(run_id)
            raise SafeStopRequested("safe stop requested")

    class StoppingRuntime(PathAwareAnalyzer):
        run_state = store
        segment_pipeline = StoppingPipeline()

    result = analyze_pending(workspace, StoppingRuntime())

    record = CatalogDatabase(workspace.database_path).list_records()[0]
    assert record.status is Status.PENDING
    assert result.failed == 0
    assert result.remaining == 1
    run = store.get_run(store.run_id_for_root(workspace.root))
    assert run is not None and run.status == "incomplete"


def test_force_batch_skips_reviewed_and_requeues_unreviewed_once(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(
        tmp_path, ("reviewed.jpg", "eligible.jpg")
    )
    analyze_pending(workspace, PathAwareAnalyzer())
    database = CatalogDatabase(workspace.database_path)
    records = {item.path.name: item for item in database.list_records()}
    reviewed_path = str(records["reviewed.jpg"].path.resolve())
    store = RunStateStore(
        workspace.database_path, excel_path=workspace.excel_path
    )
    run, _ = store.begin_run(
        root_path=workspace.root,
        video_count=0,
        image_count=2,
        total_bytes=10,
        mode=AnalysisMode.FORCE_GEMINI,
    )
    force_images = RecordingForceImages(ForceImageResult(LOCAL, None, True))
    runtime = ForceRuntime(store, force_images)

    first = analyze_pending(
        workspace,
        runtime,
        mode=AnalysisMode.FORCE_GEMINI,
        run_id=run.run_id,
        reviewed_paths={reviewed_path},
    )
    second = analyze_pending(
        workspace,
        runtime,
        mode=AnalysisMode.FORCE_GEMINI,
        run_id=run.run_id,
        reviewed_paths={reviewed_path},
    )

    assert force_images.sources == [records["eligible.jpg"].path.resolve()]
    assert first.analyzed == 1
    assert second.analyzed == 0
    reviewed = CatalogDatabase(workspace.database_path).get_record(
        records["reviewed.jpg"].id
    )
    assert reviewed is not None
    assert reviewed.description == "賀卡預覽"


def test_force_batch_saves_local_result_with_nonfatal_gemini_warning(
    tmp_path: Path,
) -> None:
    workspace = _workspace_with_media(tmp_path, ("photo.jpg",))
    store = RunStateStore(
        workspace.database_path, excel_path=workspace.excel_path
    )
    run, _ = store.begin_run(
        root_path=workspace.root,
        video_count=0,
        image_count=1,
        total_bytes=5,
        mode=AnalysisMode.FORCE_GEMINI,
    )
    force_images = RecordingForceImages(
        ForceImageResult(LOCAL, "Gemini 強化失敗:GeminiError", False)
    )

    result = analyze_pending(
        workspace,
        ForceRuntime(store, force_images),
        mode=AnalysisMode.FORCE_GEMINI,
        run_id=run.run_id,
        reviewed_paths=set(),
    )

    record = CatalogDatabase(workspace.database_path).list_records()[0]
    assert record.status is Status.ANALYZED
    assert record.description == LOCAL.description
    assert record.error == "Gemini 強化失敗:GeminiError"
    assert result.failed == 1
